#!/usr/bin/env python3
"""books-71 recon: measure the SHAPE of Michael's Cultivera exports.

Read-only. Prints sheet names, dimensions, and header rows verbatim.
No interpretation, no guessing at meaning. Rule 1.
"""
import sys
import openpyxl


def shape(path: str) -> None:
    print("=" * 78)
    print("FILE:", path)
    print("=" * 78)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print("sheet names:", wb.sheetnames)
    for ws in wb.worksheets:
        print("-" * 78)
        print(f"SHEET {ws.title!r}  max_row={ws.max_row}  max_col={ws.max_column}")
        rows = ws.iter_rows(min_row=1, max_row=6, values_only=True)
        for i, row in enumerate(rows, start=1):
            cells = ["" if c is None else str(c) for c in row]
            # trim trailing empties for readability
            while cells and cells[-1] == "":
                cells.pop()
            print(f"  row{i}: {cells}")
    wb.close()


if __name__ == "__main__":
    for p in sys.argv[1:]:
        shape(p)
